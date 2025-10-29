"""
Assistente Jurídico Inteligente — Protótipo de baixa fidelidade (Streamlit)

Esta versão adiciona:
- Landing page (Início) com cartões de features, guia de fluxo, exemplos de prompts e templates
- Navegação lateral: Início | Gerar/Editar | Exportar
- Sidebar melhorada (Sessão, Atalhos, Catálogo Normativo, Ajuda & LGPD)
- Mantém: formulário de peça, upload/edição, resumo extrativo, exportação DOCX

Fonte de requisitos (tipos de peças & referências normativas): desafio APAC/PE.
"""

from io import BytesIO
from typing import List

import streamlit as st

try:
    from docx import Document  # type: ignore
except ImportError:
    Document = (
        None  # export DOCX ficará desativado se python-docx não estiver instalado
    )


# -----------------------------------------------------------------------------
# Helpers básicos


def summarize_text(text: str, max_sentences: int = 3) -> str:
    """Resumo extrativo ingênuo (placeholder para LLM isolada)."""
    sentences = [s.strip() for s in text.replace("\n", " ").split(".") if s.strip()]
    return ". ".join(sentences[:max_sentences]) + ("." if sentences else "")


def generate_docx(content: str) -> BytesIO:
    """Gera DOCX em memória a partir de texto simples."""
    buffer = BytesIO()
    if Document is None:
        return buffer
    doc = Document()
    for paragraph in content.split("\n"):
        doc.add_paragraph(paragraph)
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def load_docx(file) -> str:
    """Extrai texto de um DOCX enviado (fallback vazio em erro/ausência do pacote)."""
    if Document is None:
        return ""
    try:
        doc = Document(file)
        return "\n".join(p.text for p in doc.paragraphs)
    except Exception:
        return ""


# -----------------------------------------------------------------------------
# Constantes derivadas do desafio (lista de peças e referências normativas)
# Requisitos oficiais do desafio — catálogos exibidos na UI (sidebar / templates).
DOCUMENT_TYPES: List[str] = [
    "Nota técnica",
    "Ofício",
    "Contratos, distratos e rescisões",
    "Aditivos contratuais (prorrogação e reajustes)",
    "Termos de cessão",
    "Convênios",
    "Termo Referencial",
    "Editais",
    "Portarias",
    "Resoluções",
    "Decretos",
    "Leis estaduais",
    "Despachos",
    "Relatórios",
    "Processos administrativos (penalidade e disciplinar)",
    "Notificações",
    "Acordos de cooperação",
    "Termos de compromisso",
]

NORMATIVE_REFERENCES: List[str] = [
    "Lei nº 14.133/2021 (Nova Lei de Licitações)",
    "Lei nº 8.666/1993 (vigente em transição)",
    "Lei nº 6.123/1968 (Estatuto dos Servidores do Estado de Pernambuco)",
    "Decreto estadual nº 52.359/2022",
    "Decreto nº 42.191/2015",
    "Decreto nº 57.002/2024",
    "Lei nº 10.406, de 10 de janeiro de 2002",
    "Lei nº 13.105, de 16 de março de 2015",
    "Constituição da República Federativa do Brasil de 1988",
    "Lei nº 14.028, de 26 de março de 2010",
    "Resoluções da ANA",
    "Boletins da Procuradoria Geral do Estado – PGE",
    "Lei nº 12.984, de 30 de dezembro de 2005",
    "Lei nº 12.334, de 20 de setembro de 2010",
]


# -----------------------------------------------------------------------------
# UI helpers (landing/feature cards/sidebar)

FEATURE_CARDS = [
    (
        "📝 Editor colaborativo",
        "Edição direta do texto com upload de .txt/.docx e exportação para .docx.",
    ),
    (
        "🧠 LLM isolada (futuro)",
        "Geração/edição guiada por prompts e padronização de linguagem jurídica.",
    ),
    (
        "⚖️ Fundamentação integrada",
        "Catálogo de leis/decretos; futuras integrações com bases jurídicas.",
    ),
    (
        "🔒 LGPD & sigilo",
        "Padrões de privacidade, anonimização e controle de acesso (IAM) a serem aplicados.",
    ),
    (
        "📑 Templates padronizados",
        "Início rápido por tipo de peça com placeholders para partes/objeto/prazos.",
    ),
    (
        "🧾 LaTeX backend (futuro)",
        "Renderização/normalização de formatação a partir de comandos de alto nível.",
    ),
]

PROMPT_EXAMPLES = [
    (
        "Ofício — abertura",
        "Gerar ofício para [órgão] comunicando [objeto] com tom formal, referenciando [norma].",
    ),
    (
        "Contrato — cláusula",
        "Inserir cláusula de reajuste com base no IPCA, periodicidade anual, conforme [Lei].",
    ),
    (
        "Relatório — sumário",
        "Criar sumário executivo com 3 parágrafos destacando [risco], [prazo] e [fundamento].",
    ),
    (
        "PAD — despacho",
        "Esboçar despacho inicial no processo disciplinar nº [0000], citando [norma] e prazo de defesa.",
    ),
]


def show_normative_references():
    st.sidebar.markdown("**Fundamentação Jurídica Integrada**")
    with st.sidebar.expander(
        "Catálogo normativo (estático neste protótipo)", expanded=False
    ):
        for ref in NORMATIVE_REFERENCES:
            st.sidebar.write(f"- {ref}")


def sidebar_layout():
    st.sidebar.markdown("### Sessão")
    if "logged_in" not in st.session_state:
        st.session_state["logged_in"] = False
    if not st.session_state["logged_in"]:
        with st.sidebar.expander("Autenticação", expanded=True):
            st.write("Demo: usuário **admin**, senha **password**.")
            u = st.text_input("Usuário", key="login_user")
            p = st.text_input("Senha", type="password", key="login_pass")
            if st.button("Entrar"):
                if u == "admin" and p == "password":
                    st.session_state["logged_in"] = True
                    st.success("Login realizado.")
                else:
                    st.error("Credenciais inválidas.")
    st.sidebar.markdown("---")
    st.sidebar.markdown("### Atalhos")
    nav_choice = st.sidebar.radio(
        "Navegação", ["Início", "Gerar/Editar", "Exportar"], index=0
    )
    st.sidebar.caption("Estado do protótipo: LLM isolada ❌  |  LaTeX backend ❌")
    show_normative_references()
    st.sidebar.markdown("---")
    with st.sidebar.expander("Ajuda & LGPD", expanded=False):
        st.write(
            "Este é um protótipo local. Em produção, será obrigatório cumprir integralmente a LGPD, "
            "com IAM, anonimização e isolamento de dados/processos."
        )
    return nav_choice


def feature_cards():
    left, right = st.columns(2)
    cards = [left, right]
    for i, (title, desc) in enumerate(FEATURE_CARDS):
        with cards[i % 2].container():
            st.markdown(f"#### {title}")
            st.write(desc)
            st.divider()


def landing_home():
    st.title("Assistente Jurídico Inteligente — APAC/PE (Protótipo)")
    st.markdown(
        "Plataforma web **SaaS-like** com editor, formulários de dados e exportação. "
        "Esta **página inicial** demonstra as capacidades previstas, mesmo sem os back-ends finais."
    )
    # métricas rápidas
    a, b, c = st.columns(3)
    a.metric("Tipos de peças", len(DOCUMENT_TYPES))
    b.metric("Referências normativas (catálogo)", len(NORMATIVE_REFERENCES))
    c.metric("Status LLM/LaTeX", "Em design", "protótipo")

    tabs = st.tabs(["Visão geral", "Como funciona", "Exemplos de prompts", "Templates"])
    with tabs[0]:
        feature_cards()
        st.info(
            "⚠️ Protótipo: integrações com bases jurídicas, LLM isolada e backend LaTeX "
            "não estão habilitadas nesta versão."
        )
    with tabs[1]:
        st.markdown("#### Fluxo recomendado")
        st.markdown(
            "- **1) Selecionar tipo de peça** e preencher dados essenciais (partes, objeto, prazos)\n"
            "- **2) Importar documento** (opcional) e **editar** no editor\n"
            "- **3) Gerar *resumo* (placeholder)** para revisão rápida\n"
            "- **4) Exportar para DOCX** (PDF e LaTeX previstos para a próxima etapa)\n"
        )
        st.button(
            "Ir para Gerar/Editar",
            key="start_generate",
            on_click=lambda: st.session_state.update(nav="Gerar/Editar"),
        )
    with tabs[2]:
        for label, example in PROMPT_EXAMPLES:
            with st.expander(f"{label}"):
                st.code(example)
    with tabs[3]:
        st.caption("Modelos baseados nos tipos de peça do desafio.")
        cols = st.columns(3)
        for i, doc in enumerate(DOCUMENT_TYPES[:9]):  # exibir 9 na grade
            with cols[i % 3]:
                st.write(f"**{doc}**")
                st.button(
                    f"Usar template",
                    key=f"tpl_{i}",
                    on_click=lambda d=doc: st.session_state.update(
                        selected_template=d, nav="Gerar/Editar"
                    ),
                )
        st.caption(
            "Use a sidebar para navegar e o formulário para preencher os campos da peça."
        )


# -----------------------------------------------------------------------------
# Telas principais


def screen_generate_edit():
    st.header("Gerar / Editar documento")
    # seleção do tipo (respeita template escolhido na landing)
    default_idx = 0
    if (
        "selected_template" in st.session_state
        and st.session_state["selected_template"] in DOCUMENT_TYPES
    ):
        default_idx = DOCUMENT_TYPES.index(st.session_state["selected_template"])
    doc_type = st.selectbox("Tipo de peça", DOCUMENT_TYPES, index=default_idx)

    with st.form("form_dados"):
        st.write("**Dados essenciais**")
        parte_autora = st.text_input("Parte Autora", placeholder="Nome da parte autora")
        parte_re = st.text_input(
            "Parte Ré/Interessada", placeholder="Nome da parte ré/interessada"
        )
        objeto = st.text_input("Objeto", placeholder="Descrição resumida do objeto")
        prazo = st.text_input("Prazos", placeholder="Data/prazo relevante")
        st.form_submit_button("Aplicar")

    st.subheader("Documento")
    uploaded_file = st.file_uploader("Upload (.txt ou .docx)", type=["txt", "docx"])
    text_content = ""
    if uploaded_file is not None:
        if uploaded_file.type == "text/plain":
            text_content = uploaded_file.read().decode("utf-8", errors="ignore")
        else:
            text_content = load_docx(uploaded_file) or "Não foi possível ler o DOCX."
    content = st.text_area(
        "Editor (conteúdo do documento)",
        height=350,
        value=text_content,
        help="Edição local no navegador (nenhum dado é enviado a serviços externos neste protótipo).",
        key="doc_content",
    )

    cols = st.columns(2)
    with cols[0]:
        if st.button("Gerar resumo (placeholder)"):
            resumo = summarize_text(st.session_state.get("doc_content", ""))
            st.info(resumo or "Nenhum conteúdo para resumir.")
    with cols[1]:
        st.caption(
            "Futuro: comandos de alto nível (ex.: “inserir cláusula X”, “padronizar linguagem Y”)."
        )


def screen_export():
    st.header("Exportar")
    content = st.session_state.get("doc_content", "")
    if not content:
        st.warning("Nada para exportar. Crie/edite um documento em **Gerar/Editar**.")
        return
    if st.button("Exportar para DOCX"):
        buf = generate_docx(content)
        if buf.getbuffer().nbytes == 0 and Document is None:
            st.error("Instale `python-docx` para habilitar a exportação.")
        else:
            st.download_button(
                "Baixar DOCX",
                data=buf,
                file_name="documento_juridico.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
    st.caption("Futuro: exportação para PDF e compilação LaTeX.")


# -----------------------------------------------------------------------------
# App


def main():
    st.set_page_config(page_title="Assistente Jurídico Inteligente", layout="wide")
    nav_choice = sidebar_layout()

    # Atalho de navegação a partir dos botões da landing/templates
    if "nav" in st.session_state and st.session_state["nav"]:
        nav_choice = st.session_state.pop("nav")

    if not st.session_state.get("logged_in", False):
        st.title("Bem-vinda(o)!")
        st.write(
            "Faça login na **sidebar** para acessar as funcionalidades do protótipo."
        )
        landing_home()  # mostra a landing mesmo sem login para dar contexto visual
        return

    if nav_choice == "Início":
        landing_home()
    elif nav_choice == "Gerar/Editar":
        screen_generate_edit()
    else:
        screen_export()


if __name__ == "__main__":
    main()
